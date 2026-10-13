from __future__ import annotations

import json
import re
import tempfile
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from app.core.config import settings
from app.services.report_builder import build_report

# Restrained technical-report palette (Word table shading uses hex without '#').
NAVY = "17324D"
BLUE = "2F6FED"
TEAL = "168C88"
INK = "1F2937"
MUTED = "596579"
PALE_BLUE = "EDF4FF"
PALE_TEAL = "EAF7F5"
PALE_GRAY = "F3F5F8"
BORDER = "D7DEE8"
WHITE = "FFFFFF"
AMBER_BG = "FFF4DF"


def xml_safe(value: Any) -> str:
    """Return XML 1.0-compatible text for python-docx/OpenXML."""
    text = str(value if value is not None else "")
    out: list[str] = []
    for ch in text:
        code = ord(ch)
        if code in (0x09, 0x0A, 0x0D) or 0x20 <= code <= 0xD7FF or 0xE000 <= code <= 0xFFFD or 0x10000 <= code <= 0x10FFFF:
            out.append(ch)
        elif code <= 0x1F:
            out.append(f"[0x{code:02X}]")
        elif 0xD800 <= code <= 0xDFFF:
            out.append("[invalid-surrogate]")
        else:
            out.append("[invalid-xml-char]")
    return "".join(out)


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_border(cell, **edges) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in edges:
            continue
        tag = "w: Diag".replace(" ", "") if edge == "diag" else f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in edges[edge]:
                element.set(qn(f"w:{key}"), str(edges[edge][key]))


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def _cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:cantSplit")
    tr_pr.append(marker)


def _set_cell_text(cell, value: Any, *, bold=False, size=8.5, color=INK, font="Arial") -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.04
    run = p.add_run(xml_safe(value))
    run.bold = bold
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)


def _style_table(table, *, header=True, font_size=8.3, widths=None) -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    if widths:
        # Set both the table grid and the individual cell preferred widths. Setting
        # only cell.width is not enough for LibreOffice/Word to preserve proportions.
        for col_idx, width in enumerate(widths):
            if col_idx < len(table.columns):
                table.columns[col_idx].width = Inches(width)
        for row in table.rows:
            for col_idx, cell in enumerate(row.cells):
                if col_idx < len(widths):
                    cell.width = Inches(widths[col_idx])
    for idx, row in enumerate(table.rows):
        _cant_split(row)
        if idx == 0 and header:
            _repeat_header(row)
        for col_idx, cell in enumerate(row.cells):
            if widths and col_idx < len(widths):
                cell.width = Inches(widths[col_idx])
            _set_cell_border(cell,
                top={"val": "single", "sz": "4", "color": BORDER},
                bottom={"val": "single", "sz": "4", "color": BORDER},
                left={"val": "single", "sz": "4", "color": BORDER},
                right={"val": "single", "sz": "4", "color": BORDER},
            )
            if idx == 0 and header:
                _set_cell_shading(cell, NAVY)
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.bold = True
                        r.font.color.rgb = RGBColor.from_string(WHITE)
                        r.font.size = Pt(font_size)
            else:
                if idx % 2 == 0:
                    _set_cell_shading(cell, "F8FAFC")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.size = Pt(font_size)


def _add_data_table(doc, headers, rows, widths=None, *, font_size=8.3, empty_text="No records were returned."):
    rows = list(rows or [])
    table = doc.add_table(rows=1, cols=len(headers))
    for i, label in enumerate(headers):
        _set_cell_text(table.rows[0].cells[i], label, bold=True, size=font_size, color=WHITE)
        _set_cell_shading(table.rows[0].cells[i], NAVY)
    if not rows:
        cells = table.add_row().cells
        cells[0].merge(cells[-1])
        _set_cell_text(cells[0], empty_text, size=font_size, color=MUTED)
    else:
        for row_data in rows:
            cells = table.add_row().cells
            for i, value in enumerate(row_data):
                _set_cell_text(cells[i], value, size=font_size)
    _style_table(table, header=True, font_size=font_size, widths=widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def _add_bullet(doc, text: Any, *, compact=True):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.14)
    p.paragraph_format.space_after = Pt(2 if compact else 5)
    p.paragraph_format.keep_together = True
    p.add_run(xml_safe(text))
    return p


def _add_numbered_method(doc, title: str, description: str):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title + "  ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(BLUE)
    p.add_run(description)
    return p


def _add_callout(doc, title: str, body: str, *, fill=PALE_BLUE, accent=BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.9)
    cell = table.cell(0, 0)
    _set_cell_shading(cell, fill)
    _set_cell_border(cell,
        left={"val": "single", "sz": "22", "color": accent},
        right={"val": "single", "sz": "3", "color": fill},
        top={"val": "single", "sz": "3", "color": fill},
        bottom={"val": "single", "sz": "3", "color": fill},
    )
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    rr = p.add_run(xml_safe(title))
    rr.bold = True
    rr.font.size = Pt(10)
    rr.font.color.rgb = RGBColor.from_string(accent)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(4)
    p2.add_run(xml_safe(body))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def _add_metric_cards(doc, cards):
    table = doc.add_table(rows=1, cols=len(cards))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, (label, value, hint, accent) in enumerate(cards):
        cell = table.cell(0, idx)
        cell.width = Inches(6.9 / max(1, len(cards)))
        _set_cell_shading(cell, PALE_GRAY)
        _set_cell_border(cell,
            top={"val": "single", "sz": "8", "color": accent},
            bottom={"val": "single", "sz": "4", "color": BORDER},
            left={"val": "single", "sz": "4", "color": BORDER},
            right={"val": "single", "sz": "4", "color": BORDER},
        )
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(xml_safe(value))
        r.bold = True
        r.font.size = Pt(16)
        r.font.color.rgb = RGBColor.from_string(accent)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r = p2.add_run(xml_safe(label))
        r.bold = True
        r.font.size = Pt(8.2)
        p3 = cell.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.paragraph_format.space_after = Pt(5)
        r = p3.add_run(xml_safe(hint))
        r.font.size = Pt(7.2)
        r.font.color.rgb = RGBColor.from_string(MUTED)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _add_page_field(paragraph, instruction: str):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def _friendly_event(event: str) -> str:
    known = {
        "login": "User login",
        "file_open": "Confidential report opened",
        "usb_connected": "USB device connected",
        "archive_created": "ZIP archive created",
        "delete": "Original report deleted",
        "history_clear": "Shell history clear recorded",
        "usb_disconnected": "USB device disconnected",
    }
    return known.get((event or "").lower(), (event or "Event").replace("_", " ").strip().title())


def _event_time(event: dict[str, Any]) -> str:
    raw = event.get("timestamp") or event.get("local_time") or ""
    try:
        dt = datetime.fromisoformat(str(raw).replace("Z", "+00:00"))
        zone = dt.strftime("%z")
        zone = f"UTC{zone[:3]}:{zone[3:]}" if zone else "local time"
        return f"{dt.strftime('%H:%M:%S')} {zone}"
    except Exception:
        return str(raw or "Time not recorded")


def _timestamp_sort_key(event: dict[str, Any]):
    raw = event.get("timestamp") or event.get("local_time") or ""
    try:
        return (0, datetime.fromisoformat(str(raw).replace("Z", "+00:00")).timestamp())
    except Exception:
        return (1, int(event.get("sequence") or 0))


def _short_detail(detail: Any, limit=105) -> str:
    s = re.sub(r"\s+", " ", xml_safe(detail)).strip()
    s = re.sub(r"\bserial=[^\s]+", "", s, flags=re.I).strip()
    return s if len(s) <= limit else s[: limit - 1].rstrip() + "…"


def _make_timeline_image(events: list[dict[str, Any]], output_path: Path) -> bool:
    """Create a clean vertical event diagram from timestamped events only."""
    chron = [e for e in events if e.get("timestamp") or e.get("local_time")]
    chron = sorted(chron, key=_timestamp_sort_key)
    if not chron:
        return False
    try:
        from PIL import Image, ImageDraw, ImageFont
    except ImportError:
        return False

    width = 1480
    row_h = 142
    top = 65
    height = top + len(chron) * row_h + 35
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_paths = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "C:/Windows/Fonts/arial.ttf",
    ]
    bold_paths = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
        "C:/Windows/Fonts/arialbd.ttf",
    ]
    def load_font(paths, size):
        for path in paths:
            try:
                return ImageFont.truetype(path, size)
            except Exception:
                continue
        return ImageFont.load_default()
    f_time = load_font(bold_paths, 28)
    f_event = load_font(bold_paths, 30)
    f_detail = load_font(font_paths, 23)
    f_small = load_font(font_paths, 20)

    draw.rounded_rectangle((24, 20, width - 24, height - 20), radius=26, fill="#F5F8FC", outline="#D8E1ED", width=2)
    x_line = 112
    y_centers = [top + i * row_h + 42 for i in range(len(chron))]
    if len(y_centers) > 1:
        draw.line((x_line, y_centers[0], x_line, y_centers[-1]), fill="#B5C7E2", width=8)
    for idx, event in enumerate(chron):
        cy = y_centers[idx]
        color = TEAL if (event.get("event") or "").lower() in {"usb_connected", "archive_created"} else BLUE
        # Use Python style hex constants; Word uses unprefixed hex for table fills.
        color = "#168C88" if color == TEAL else "#2F6FED"
        draw.ellipse((x_line - 22, cy - 22, x_line + 22, cy + 22), fill=color, outline="white", width=5)
        time_label = _event_time(event)
        title = _friendly_event(event.get("event") or "event")
        detail = _short_detail(event.get("details") or event.get("source") or "Event recorded in evidence.", 100)
        draw.text((175, cy - 45), time_label, fill="#526175", font=f_time)
        draw.text((520, cy - 45), title, fill="#17324D", font=f_event)
        draw.text((520, cy + 2), detail, fill="#526175", font=f_detail)
        src = event.get("source") or ""
        if src:
            draw.text((520, cy + 33), src, fill="#748197", font=f_small)
    image.save(output_path, format="PNG", optimize=True)
    return True


def _plain_text_preview(content: Any, path: str = "", limit=1400) -> str:
    text = str(content or "")
    lower_path = path.lower()
    if lower_path.endswith((".zip", ".gz", ".7z", ".rar", ".png", ".jpg", ".jpeg", ".bin")) or text.startswith("PK\x03\x04"):
        return ""
    if text.startswith("%PDF"):
        # Keep a useful textual string, not the PDF object/binary dump.
        matched = re.findall(r"\(([^()]{4,240})\)\s*Tj", text)
        return "\n".join(matched[:8])[:limit]
    if not text:
        return ""
    printable = sum(1 for ch in text if ch.isprintable() or ch in "\r\n\t") / max(1, len(text))
    if printable < 0.88:
        return ""
    cleaned = "".join(ch if ch.isprintable() or ch in "\r\n\t" else " " for ch in text)
    cleaned = re.sub(r"\n{3,}", "\n\n", cleaned).strip()
    return cleaned[:limit]


def _add_running_header_footer(doc):
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.paragraph_format.space_after = Pt(2)
        r = header.add_run("EXT4 FORENSICS WORKBENCH  /  CYBER FORENSIC EXAMINATION")
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor.from_string(MUTED)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.paragraph_format.space_before = Pt(3)
        r = footer.add_run("CONFIDENTIAL · EXAMINER REVIEW REQUIRED     |     PAGE ")
        r.font.name = "Arial"
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor.from_string(MUTED)
        _add_page_field(footer, "PAGE")
        r = footer.add_run(" OF ")
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor.from_string(MUTED)
        _add_page_field(footer, "NUMPAGES")


def _setup_document() -> Document:
    doc = Document()
    sec = doc.sections[0]
    # Explicit A4 portrait configuration.
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.65)
    sec.bottom_margin = Cm(1.6)
    sec.left_margin = Cm(1.7)
    sec.right_margin = Cm(1.7)
    sec.header_distance = Cm(0.8)
    sec.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for style_name, size, color in [
        ("Title", 25, NAVY), ("Heading 1", 15, NAVY), ("Heading 2", 11.2, BLUE), ("Heading 3", 9.8, TEAL)
    ]:
        s = styles[style_name]
        s.font.name = "Arial"
        s.font.size = Pt(size)
        s.font.bold = style_name != "Title" or True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(9 if style_name == "Heading 1" else 5)
        s.paragraph_format.space_after = Pt(4)
        s.paragraph_format.keep_with_next = True
    _add_running_header_footer(doc)
    return doc


def _add_title(doc, case_number: str, case_title: str, scan_id: str):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("CYBER FORENSIC\nEXAMINATION REPORT")
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor.from_string(NAVY)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    r = p.add_run("CONFIDENTIAL FORENSIC WORK PRODUCT  ·  EXAMINER REVIEW REQUIRED")
    r.bold = True
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    metadata = [
        ("CASE NUMBER", case_number or "Not supplied"),
        ("CASE TITLE", case_title or "Cyber forensic examination"),
        ("SCAN REFERENCE", (scan_id or "Not supplied")[:18]),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    for idx, (label, value) in enumerate(metadata):
        cell = table.cell(0, idx)
        cell.width = Inches(6.9 / 3)
        _set_cell_shading(cell, PALE_BLUE if idx != 1 else PALE_TEAL)
        _set_cell_border(cell,
            top={"val": "single", "sz": "5", "color": BORDER},
            bottom={"val": "single", "sz": "5", "color": BORDER},
            left={"val": "single", "sz": "5", "color": BORDER},
            right={"val": "single", "sz": "5", "color": BORDER},
        )
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(label)
        r.bold = True
        r.font.size = Pt(7.2)
        r.font.color.rgb = RGBColor.from_string(MUTED)
        p2 = cell.add_paragraph()
        p2.paragraph_format.space_after = Pt(4)
        r = p2.add_run(xml_safe(value))
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor.from_string(NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _section(doc, title: str):
    p = doc.add_heading(title, level=1)
    p.paragraph_format.keep_with_next = True
    return p


def _technical_method(doc):
    _section(doc, "Examination method")
    p = doc.add_paragraph(
        "The examination uses a read-only workflow against the supplied disk image. The scanner does not mount the source volume or write to the evidence image. "
        "The following stages describe what the workbench actually performs; the scope is limited to the supplied image and supported EXT4 structures."
    )
    p.paragraph_format.keep_together = True
    _add_numbered_method(doc, "1. Intake and integrity baseline.",
        "The image path is validated and the image is streamed through SHA-256 hashing before analysis. Its byte size and pre-scan digest are recorded with the case/evidence identifiers. "
        "The hash provides a reproducible integrity reference; it does not independently prove who created or acquired the image.")
    _add_numbered_method(doc, "2. Filesystem location and metadata scan.",
        "The scanner checks the image start and detected partition offsets for an EXT4 superblock, then reads filesystem geometry such as UUID, volume label, block size, inode size and group counts. "
        "It walks block-group inode tables using bounded I/O workers and identifies deleted regular files, directories and symbolic links where inode metadata indicates deletion.")
    _add_numbered_method(doc, "3. Name/path reconstruction and cross-checking.",
        "Surviving directory entries and directory slack are examined to associate inode numbers with candidate names and parent directories. This reconstruction is metadata-driven and may be incomplete or ambiguous. "
        "Where available, pytsk3 independently opens the filesystem and traverses directory entries to extract targeted live artifacts. Targeted JBD2 journal inspection supplements the analysis; it is not a full journal replay.")
    _add_numbered_method(doc, "4. Artifact and timeline extraction.",
        "Selected files such as activity logs and shell history are read through filesystem directory entries, subject to configured size limits. Activity-log records with timestamps become time-based events. "
        "Shell-history commands without embedded timestamps are retained as sequence evidence only; the report does not assign them invented clock times. Each extracted artifact is hashed and its source path, inode and filesystem timestamps are recorded when available.")
    _add_numbered_method(doc, "5. Deleted-file recovery and validation.",
        "For a selected deleted-file candidate, the recovery routine reads the inode and follows retained EXT4 extents or supported legacy block pointers, writing reconstructed bytes to the configured recovery directory rather than altering the source image. "
        "The tool compares bytes written with the inode's logical size, flags incomplete or unwritten extents, inspects leading bytes for a recognizable file signature and computes a SHA-256 digest of the recovered output. A candidate is not described as fully recovered unless the recovery checks support that status.")
    _add_numbered_method(doc, "6. Post-scan integrity and audit trail.",
        "After scanning, the source image is hashed again. The pre- and post-scan SHA-256 values are compared and scan/hash events are written to a local hash-chained JSONL integrity ledger. "
        "The ledger chain is verified before report generation. Matching hashes support that the image bytes did not change during this run, subject to the reliability of the measurement environment.")
    _add_callout(doc, "Recovery scope note",
        "This workbench primarily performs metadata-guided recovery using surviving inode and data-block references. Signature detection helps identify/validate file type, but this report must not be read as a claim that a comprehensive raw-block carving pass or a full JBD2 replay was performed.",
        fill=AMBER_BG, accent="A96708")


def _compact_timestamp(value: Any, *, include_zone=False) -> str:
    """Format timestamps for tables without wrapping ISO strings awkwardly."""
    raw = str(value or "")
    if not raw or raw in {"—", "-"}:
        return raw or "—"
    try:
        dt = datetime.fromisoformat(raw.replace("Z", "+00:00"))
        first = dt.strftime("%d %b %Y")
        second = dt.strftime("%H:%M:%S")
        if include_zone:
            zone = dt.strftime("%z")
            if zone:
                second += f" UTC{zone[:3]}:{zone[3:]}"
            else:
                second += " (timezone not recorded)"
        return first + "\n" + second
    except Exception:
        return raw


def _pretty_hash(value: Any) -> str:
    digest = str(value or "")
    if re.fullmatch(r"[A-Fa-f0-9]{64}", digest):
        return digest[:32] + "\n" + digest[32:]
    return digest or "Not available"

def _format_hash(value: Any) -> str:
    return _pretty_hash(value)


def _event_details(event: dict[str, Any]) -> str:
    details = str(event.get("details") or "—")
    return details.replace("action=", "action=")


def render_report(report: dict[str, Any], output_path: str | Path, scan_id: str = "") -> str:
    """Render a report dictionary to an A4 DOCX. Exposed for tests and offline report assembly."""
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = _setup_document()

    case = report.get("case", {}) or {}
    evidence = report.get("evidence", {}) or {}
    ai = report.get("ai_report", {}) or {}
    generation = ai.get("generation", {}) or {}
    fs = report.get("filesystem_metadata", {}) or {}
    deleted = report.get("deleted_files", []) or []
    timeline = report.get("timeline", []) or []
    artifacts = report.get("live_artifacts", []) or []
    ledger = report.get("integrity_ledger", {}) or {}

    if not scan_id:
        scan_id = str(report.get("scan_id") or "")
    _add_title(doc, case.get("case_number") or "Not supplied", case.get("name") or "Cyber forensic examination", scan_id)

    _section(doc, "Executive summary")
    summary = ai.get("executive_summary") or (
        "The supplied evidence image was examined for EXT4 filesystem metadata, deleted-file candidates and relevant activity artifacts. "
        "Refer to the findings, event timeline and limitations below. Interpretations remain subject to examiner review."
    )
    narrative_label = "Overall assessment"
    if generation.get("provider") == "Groq":
        narrative_label += " · Groq-assisted narrative"
    elif generation.get("status") in {"fallback", "not_configured"}:
        narrative_label += " · Local fallback narrative"
    _add_callout(doc, narrative_label, summary, fill=PALE_BLUE, accent=BLUE)
    complete_recoveries = sum(1 for item in deleted if str(item.get("recovery_status") or "").lower() == "complete")
    _add_metric_cards(doc, [
        ("DELETED-FILE CANDIDATES", str(len(deleted)), "Metadata candidates identified", BLUE),
        ("IMAGE INTEGRITY", "MATCH" if evidence.get("integrity_match") else "UNVERIFIED", "Pre / post SHA-256 comparison", TEAL if evidence.get("integrity_match") else "A96708"),
        ("COMPLETE RECOVERIES", str(complete_recoveries), "Validated complete outputs", TEAL if complete_recoveries else "A96708"),
    ])

    objectives = ai.get("investigation_objectives")
    if objectives:
        doc.add_heading("Examination objectives", level=2)
        doc.add_paragraph(xml_safe(objectives))
    if ai.get("key_findings"):
        doc.add_heading("Key findings", level=2)
        for finding in ai["key_findings"]:
            _add_bullet(doc, finding)
    else:
        facts = (report.get("lab02_deliverables", {}).get("8_facts_and_hypotheses", {}) or {}).get("facts", [])
        for finding in facts[:6]:
            _add_bullet(doc, finding)

    _section(doc, "1. Case and evidence")
    compact_rows = [
        ("Evidence item", evidence.get("evidence_item_number") or "Not assigned"),
        ("Evidence type", evidence.get("evidence_type") or "Not supplied"),
        ("Evidence description", evidence.get("description") or ""),
        ("Source device / asset", evidence.get("source_device") or ""),
    ]
    compact_rows = [(k, v) for k, v in compact_rows if str(v or "").strip()]
    _add_data_table(doc, ["Field", "Recorded value"], compact_rows, widths=[2.0, 4.9], font_size=8.6,
                    empty_text="Only case number and case title were provided.")

    _section(doc, "2. Evidence integrity")
    integrity_rows = [
        ("Image path", evidence.get("path") or "Not recorded"),
        ("Image size", f'{int(evidence.get("size_bytes") or 0):,} bytes' if evidence.get("size_bytes") is not None else "Not recorded"),
        ("Pre-scan SHA-256", _format_hash(evidence.get("pre_scan_sha256") or evidence.get("sha256"))),
        ("Post-scan SHA-256", _format_hash(evidence.get("post_scan_sha256"))),
        ("Comparison", "MATCH — image hashes agree" if evidence.get("integrity_match") else "NOT VERIFIED — hashes missing or different"),
    ]
    _add_data_table(doc, ["Check", "Result"], integrity_rows, widths=[2.0, 4.9], font_size=8.1)

    _technical_method(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    policy = (report.get("tooling", {}) or {}).get("resource_policy", {}) or {}
    if policy:
        p.add_run("Execution profile: ").bold = True
        p.add_run(xml_safe(f'policy={policy.get("policy", "default")}; I/O workers={policy.get("io_workers", "—")}; hash workers={policy.get("hash_workers", "—")}'))

    _section(doc, "3. Filesystem findings")
    fs_rows = [
        ("Filesystem", report.get("filesystem") or "Not extracted"),
        ("Volume label", report.get("volume_label") or "Not extracted"),
        ("Filesystem UUID", fs.get("uuid") or "—"),
        ("Block size", f'{fs.get("block_size"):,} bytes' if fs.get("block_size") else "—"),
        ("Inode count", f'{fs.get("inodes_count"):,}' if fs.get("inodes_count") is not None else "—"),
        ("Block count", f'{fs.get("blocks_count"):,}' if fs.get("blocks_count") is not None else "—"),
        ("Filesystem offset", f'{fs.get("filesystem_offset"):,} bytes' if fs.get("filesystem_offset") is not None else "—"),
    ]
    _add_data_table(doc, ["Property", "Observed value"], fs_rows, widths=[2.0, 4.9], font_size=8.4)

    _section(doc, "4. Deleted-file findings")
    deleted_rows = []
    for f in deleted:
        deleted_rows.append((
            f.get("inode", "—"), f.get("original_name") or "Unnamed candidate",
            f.get("original_path") or "Path not reconstructed",
            f'{f.get("size_bytes"):,} B' if isinstance(f.get("size_bytes"), (int, float)) else "—",
            _compact_timestamp(f.get("deletion_time")), (f.get("recovery_status") or "candidate").replace("_", " ").title(),
            f.get("validation_status") or "Not validated",
        ))
    _add_data_table(doc, ["Inode", "Name", "Original path", "Size", "Deletion time (UTC)", "Recovery", "Validation"],
                    deleted_rows, widths=[0.48, 1.1, 2.0, 0.55, 1.0, 0.75, 1.0], font_size=7.1)
    if deleted:
        recovered_records = [f for f in deleted if f.get("recovered_path") or f.get("recovered_size_bytes") is not None or f.get("recovered_sha256")]
        if recovered_records:
            doc.add_heading("Recovered-output validation", level=2)
            recovered_rows = [(f.get("inode", "—"), f.get("recovered_path") or "Path not recorded", f.get("recovered_size_bytes") if f.get("recovered_size_bytes") is not None else "—", _pretty_hash(f.get("recovered_sha256")) or "—") for f in recovered_records]
            _add_data_table(doc, ["Inode", "Recovered output", "Bytes written", "Recovered SHA-256"], recovered_rows,
                            widths=[0.55, 2.3, 0.8, 3.2], font_size=7.4)
        else:
            p = doc.add_paragraph("No recovered output path, byte count or recovered SHA-256 is recorded for the listed deleted-file candidates.")
            p.paragraph_format.space_after = Pt(4)

    _section(doc, "5. Timeline assessment")
    timeline_analysis = ai.get("timeline_analysis") or "A time-ordered assessment is based on timestamped activity-log records. Shell history entries without timestamps are treated as sequence evidence and are not assigned clock times."
    timeline_text = doc.add_paragraph(xml_safe(timeline_analysis))
    timeline_text.paragraph_format.keep_with_next = True
    with tempfile.TemporaryDirectory(prefix="ext4-timeline-") as temp_dir:
        timeline_png = Path(temp_dir) / "timeline.png"
        if _make_timeline_image(timeline, timeline_png):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            p.add_run().add_picture(str(timeline_png), width=Inches(5.95))
            cap = doc.add_paragraph("Figure 1. Timestamped events extracted from the evidence image. Times retain the timezone offset recorded by the source.")
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.space_after = Pt(5)
            for r in cap.runs:
                r.italic = True
                r.font.size = Pt(7.5)
                r.font.color.rgb = RGBColor.from_string(MUTED)

    timed = [e for e in timeline if e.get("timestamp") or e.get("local_time")]
    timed.sort(key=_timestamp_sort_key)
    event_rows = [(_event_time(e), _friendly_event(e.get("event") or "event"), Path(str(e.get("source") or "—")).name, _event_details(e)) for e in timed]
    _add_data_table(doc, ["Recorded time", "Event", "Source", "Details"], event_rows,
                    widths=[1.25, 1.45, 1.05, 3.1], font_size=7.0)

    sequence_events = [e for e in timeline if not (e.get("timestamp") or e.get("local_time"))]
    if sequence_events:
        doc.add_heading("Sequence-only shell history evidence", level=2)
        seq_rows = []
        for e in sorted(sequence_events, key=lambda x: int(x.get("sequence") or 0)):
            seq_rows.append((f'Command #{e.get("sequence", "—")}', e.get("event", "shell_command").replace("_", " ").title(), e.get("details") or "—", "Timestamp not present"))
        _add_data_table(doc, ["Order", "Classification", "Recorded command", "Timing basis"], seq_rows,
                        widths=[0.8, 1.2, 3.8, 1.05], font_size=7.4)

    inventory_heading = _section(doc, "6. Live artifact inventory")
    inventory_heading.paragraph_format.page_break_before = True
    artifact_rows = []
    for a in artifacts:
        artifact_rows.append((
            a.get("path") or "—", a.get("inode", "—"),
            f'{a.get("size", a.get("size_bytes", 0)):,} B' if isinstance(a.get("size", a.get("size_bytes", 0)), (int, float)) else "—",
            _compact_timestamp(a.get("mtime")), _pretty_hash(a.get("content_sha256")) or "—",
        ))
    _add_data_table(doc, ["Path", "Inode", "Size", "Mtime", "Content SHA-256"], artifact_rows,
                    widths=[1.75, 0.45, 0.55, 1.15, 2.95], font_size=6.8)

    _section(doc, "7. Supporting artifact excerpts")
    shown = 0
    for a in artifacts:
        path = str(a.get("path") or "")
        content = _plain_text_preview(a.get("content"), path=path, limit=1000)
        if not content:
            continue
        base = Path(path).name or path
        doc.add_heading(base, level=2)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.right_indent = Inches(0.1)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_together = True
        r = p.add_run(xml_safe(content))
        r.font.name = "Consolas"
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor.from_string(INK)
        shown += 1
        if shown >= 6:
            break
    if shown == 0:
        doc.add_paragraph("No text artifact previews were included in this report. See artifact metadata and the machine-readable report for further detail.")

    facts_heading = _section(doc, "8. Facts, interpretation and alternative explanations")
    facts_heading.paragraph_format.page_break_before = True
    deliverables = report.get("lab02_deliverables", {}) or {}
    fh = deliverables.get("8_facts_and_hypotheses", {}) or {}
    facts = fh.get("facts", [])
    hypotheses = fh.get("hypotheses", [])
    if facts:
        doc.add_heading("Observed facts", level=2)
        for fact in facts:
            _add_bullet(doc, fact)
    interpretation = ai.get("interpretation")
    if interpretation:
        _add_callout(doc, "Interpretation — not a direct observation", interpretation, fill=PALE_TEAL, accent=TEAL)
    if hypotheses:
        doc.add_heading("Hypotheses / interpretation requiring corroboration", level=2)
        for hypothesis in hypotheses:
            _add_bullet(doc, hypothesis)
    _add_callout(doc, "Transfer conclusion",
        "The temporal proximity between USB connection and archive creation may be consistent with a transfer workflow, but the supplied image does not establish that the ZIP archive was successfully written to the USB device. A USB image or other independent corroboration is required.",
        fill=AMBER_BG, accent="A96708")

    _section(doc, "9. Integrity ledger")
    _add_callout(doc, "Ledger verification", ledger.get("message") or "Ledger verification result was not supplied.",
                 fill=PALE_TEAL if ledger.get("verified") else AMBER_BG, accent=TEAL if ledger.get("verified") else "A96708")
    integrity_events = report.get("integrity_events", []) or []
    if integrity_events:
        # Keep the report audit-friendly without flooding it with duplicate run records.
        display = integrity_events[-8:]
        ledger_rows = [(_compact_timestamp(e.get("created_at"), include_zone=True), e.get("event_type") or "—", e.get("algorithm") or "—", _pretty_hash(e.get("value")) if e.get("value") else (e.get("details") or "—")) for e in display]
        _add_data_table(doc, ["Recorded at", "Event", "Algorithm", "Value / details"], ledger_rows,
                        widths=[1.4, 1.3, 0.8, 3.35], font_size=7.0)
        if len(integrity_events) > len(display):
            p = doc.add_paragraph(f'{len(integrity_events)} integrity records exist for this case; the {len(display)} most recent records are shown here. The local chained ledger is the authoritative complete audit log.')
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs:
                r.italic = True
                r.font.size = Pt(7.5)
                r.font.color.rgb = RGBColor.from_string(MUTED)

    limitations_heading = _section(doc, "10. Limitations")
    limitations_heading.paragraph_format.page_break_before = True
    limitations = list(report.get("limitations", []) or [])
    for item in ai.get("limitations", []) or []:
        if item not in limitations:
            limitations.append(item)
    # Deduplicate exact and near-duplicate limitation strings to avoid boilerplate repetition.
    seen_limitations = set()
    for item in limitations:
        key = re.sub(r"\W+", " ", str(item).lower()).strip()
        if key and key not in seen_limitations:
            _add_bullet(doc, item)
            seen_limitations.add(key)

    _section(doc, "11. Conclusion")
    doc.add_paragraph(xml_safe(ai.get("conclusion") or
        "The findings and conclusions in this report are limited to artifacts and metadata identified in the supplied disk image. "
        "The examiner should review this report against the underlying evidence and retain the image, hashes and complete audit ledger for independent verification."))

    recommendations = ai.get("recommendations", []) or []
    if recommendations:
        _section(doc, "Recommended next steps")
        for item in recommendations:
            _add_bullet(doc, item)

    # A compact traceability matrix retains the Lab #02 deliverables without duplicating raw JSON.
    if deliverables:
        _section(doc, "Appendix A. Lab #02 deliverable traceability")
        conf = deliverables.get("6_confidential_document_recovered")
        lab_rows = [
            ("Image SHA-256", "Recorded in Evidence integrity."),
            ("Filesystem and volume label", f'{deliverables.get("2_filesystem_and_volume_label", {}).get("filesystem") or "Not extracted"} / {deliverables.get("2_filesystem_and_volume_label", {}).get("volume_label") or "Not extracted"}'),
            ("Deleted files and inode numbers", f'{len(deliverables.get("3_deleted_files_and_inodes", []) or [])} candidate(s) listed.'),
            ("USB connection", "See timeline and event source." if deliverables.get("4_usb_connected") else "No connection event extracted."),
            ("Archive creation", "See timeline and event source." if deliverables.get("5_archive_created") else "No archive event extracted."),
            ("Confidential PDF recovery", "Candidate(s) listed; see recovery status and validation." if conf else "No validated deleted-PDF recovery result supplied."),
            ("Metadata recovery versus carving", "Explained in the examination method and recovery scope note."),
            ("Facts versus hypotheses", "Separated in Section 8."),
        ]
        _add_data_table(doc, ["Deliverable", "Where addressed / outcome"], lab_rows, widths=[2.25, 4.65], font_size=7.7)

    doc.core_properties.title = "Cyber Forensic Examination Report"
    doc.core_properties.subject = "Forensic examination of supplied EXT4 disk image"
    doc.core_properties.keywords = "cyber forensics, EXT4, integrity, evidence timeline"
    doc.core_properties.comments = "Confidential forensic work product. Examiner review required."
    doc.save(out)
    return str(out)


def generate_docx(scan_id, output_path=None):
    report = build_report(scan_id)
    out = Path(output_path or settings.report_dir / f"ext4-report-{scan_id[:8]}.docx")
    return render_report(report, out, scan_id=scan_id)
